"""销项发票导入（销项开票在外部系统执行）与红字发票处理。

定位：发票易不做出票，只**导入已开票**用于 ① 与收款回单对账核销 ② 销项归集与报表。

支持的导入方式：
1. **文件解析**：数电票 XML / OFD / PDF 走现有分级解析链路（原件归档）
2. **清单批量**：开票系统导出的 CSV/Excel（号码/日期/购买方/金额/税额/类型/原发票号）

红字发票处理（本次核心）：
- 红票 `red_flag=True`，**不进入待收款**；报表按"客户净额 = Σ蓝票 − Σ红票"冲减
- **自动关联原蓝票**：数电红票票面含原发票号码 → 解析后自动挂 `original_invoice_id`
- **人工补关联**：未自动关联的红票由财务在详情页指定原蓝票（`link_red_invoice`）
- 与回单的关联：红票对应银行流水通常是**退款（我方付款）**，配对走"付"方向

幂等：发票号码全局唯一（一张票不会既是进项又是销项）——重复导入按 skipped 报告。
"""
import csv
import io
import logging
from decimal import Decimal, InvalidOperation

from sqlalchemy.orm import Session

from invoicing.audit import write_audit
from invoicing.models import Invoice, User
from invoicing.models.fields import utcnow

logger = logging.getLogger(__name__)

# 清单表头别名（开票系统导出字段名不统一，做兼容映射）
_LIST_FIELD_ALIASES = {
    "invoice_number": ("发票号码", "发票号", "invoice_number", "InvoiceNumber"),
    "issue_date": ("开票日期", "日期", "issue_date", "IssueDate"),
    "buyer_name": ("购买方名称", "购方名称", "客户名称", "buyer_name", "BuyerName"),
    "buyer_tax_id": ("购买方税号", "购方税号", "buyer_tax_id", "BuyerTaxID"),
    "amount_without_tax": ("金额", "不含税金额", "amount_without_tax"),
    "tax_amount": ("税额", "tax_amount", "TaxAmount"),
    "total_amount": ("价税合计", "含税金额", "total_amount", "TotalAmount"),
    "invoice_type": ("发票类型", "票种", "invoice_type"),
    "original_invoice_number": (
        "原发票号码", "对应蓝字发票号码", "原蓝票号码", "original_invoice_number",
    ),
}

# 原发票号码在 XML 中的常见标签（数电红字票版式不一，做尽力提取）
_ORIGINAL_KEYS = ("OriginalInvoiceNumber", "原发票号码", "OriginalInvoiceNo", "对应蓝字发票号码")


def _extract_original_number(data: bytes, text: str | None) -> str | None:
    """从 XML/文本中尽力提取"原发票号码"（红字票自动关联原蓝票的关键）。"""
    import re

    haystack = ""
    try:
        haystack = data.decode("utf-8", errors="ignore")
    except Exception:
        haystack = ""
    for key in _ORIGINAL_KEYS:
        m = re.search(rf"<[^>]*{key}[^>]*>\s*([0-9A-Za-z]{{8,}})", haystack)
        if m:
            return m.group(1)
    if text:
        m = re.search(r"(?:原发票号码|对应蓝字发票号码)[:：]?\s*([0-9A-Za-z]{8,})", text)
        if m:
            return m.group(1)
    return None


def _link_original(db: Session, inv: Invoice, original_number: str | None) -> None:
    """红票自动关联原蓝票（号码命中才关联；未命中留空待人工补）。"""
    if not inv.red_flag or not original_number:
        return
    original = (
        db.query(Invoice)
        .filter(Invoice.invoice_number == original_number, Invoice.id != inv.id)
        .first()
    )
    if original is not None:
        inv.original_invoice_id = original.id


def import_sales_files(
    db: Session, user: User, files: list[tuple[str, bytes]]
) -> list[dict]:
    """文件解析导入销项票（XML/OFD/PDF）：逐文件返回 imported/skipped/error，互不影响。"""
    from invoicing.fetch.filters import classify_attachment
    from invoicing.parse.router import parse_file
    from invoicing.parse.red_flag import detect_red_invoice
    from invoicing.storage import get_storage

    results = []
    for filename, data in files:
        try:
            kind = classify_attachment(filename, "", data)
            if kind not in ("XML", "OFD", "PDF"):
                results.append({"file": filename, "status": "error",
                                "error": f"不支持的格式: {kind or '未知'}（仅 XML/OFD/PDF）"})
                continue
            outcome = parse_file(kind, data)
            parsed = outcome.parsed
            if parsed is None or not parsed.invoice_number:
                results.append({"file": filename, "status": "error",
                                "error": "未解析出发票号码，无法作为销项票导入"})
                continue
            if db.query(Invoice).filter(Invoice.invoice_number == parsed.invoice_number).first():
                results.append({"file": filename, "status": "skipped",
                                "error": "该发票号码已存在"})
                continue

            storage = get_storage()
            key = f"tenant-default/sales/{parsed.invoice_number}-{filename}"
            storage.put(key, data, "application/octet-stream")
            inv = Invoice(
                file_url=key, file_type=kind, invoice_direction="output",
                invoice_number=parsed.invoice_number, invoice_code=parsed.invoice_code,
                issue_date=parsed.issue_date, total_amount=parsed.total_amount,
                amount_without_tax=parsed.amount_without_tax, tax_amount=parsed.tax_amount,
                total_amount_cn=parsed.total_amount_cn,
                seller_name=parsed.seller_name or "", buyer_name=parsed.buyer_name,
                buyer_tax_id=parsed.buyer_tax_id, seller_tax_id=parsed.seller_tax_id,
                invoice_type=parsed.invoice_type, parse_source=parsed.parse_source,
                confidence_score=parsed.confidence_score,
                status="parsed", verify_status="pending",
                # 红字判定：优先发票类型；XML 文本兜底（数电红票的 InvoiceType
                # 未必被解析器提取，但票面文本一定含"红字"字样）
                red_flag=detect_red_invoice(
                    data.decode("utf-8", errors="ignore")[:4000] if kind == "XML" else None,
                    parsed.invoice_type,
                ),
                user_id=user.id,
            )
            db.add(inv)
            db.flush()
            _link_original(db, inv, _extract_original_number(data, None))
            write_audit(
                db, action="SALES_IMPORT", user_id=user.id, invoice_id=inv.id, channel="web",
                detail={"file": filename, "invoice_number": inv.invoice_number,
                        "red_flag": inv.red_flag,
                        "original_invoice_id": inv.original_invoice_id},
            )
            db.commit()
            results.append({"file": filename, "status": "imported", "invoice_id": inv.id,
                            "invoice_number": inv.invoice_number, "red_flag": inv.red_flag})
        except Exception as exc:  # 单个文件失败不影响其余
            db.rollback()
            logger.exception("销项文件导入失败 file=%s", filename)
            results.append({"file": filename, "status": "error", "error": str(exc)})
    return results


def _pick(row: dict, field: str):
    for alias in _LIST_FIELD_ALIASES[field]:
        if alias in row and str(row[alias]).strip() != "":
            return str(row[alias]).strip()
    return None


def _to_decimal(raw: str | None) -> Decimal | None:
    if raw is None:
        return None
    try:
        return Decimal(raw.replace(",", "").replace("￥", "").strip())
    except (InvalidOperation, AttributeError):
        return None


def import_sales_list(
    db: Session, user: User, data: bytes, filename: str = "sales.csv"
) -> dict:
    """清单批量导入（CSV/Excel）：逐行入库，重复号码跳过；红票按"原发票号码"关联蓝票。"""
    from invoicing.parse.red_flag import detect_red_invoice

    rows: list[dict] = []
    if filename.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.active
        iterator = ws.iter_rows(values_only=True)
        header = [str(c).strip() if c is not None else "" for c in next(iterator)]
        for values in iterator:
            rows.append({header[i]: values[i] for i in range(min(len(header), len(values)))})
    else:
        text = data.decode("utf-8-sig", errors="ignore")
        rows = list(csv.DictReader(io.StringIO(text)))

    imported = skipped = errors = 0
    details: list[dict] = []
    for row in rows:
        number = _pick(row, "invoice_number")
        if not number:
            errors += 1
            details.append({"status": "error", "error": "缺少发票号码"})
            continue
        if db.query(Invoice).filter(Invoice.invoice_number == number).first():
            skipped += 1
            details.append({"invoice_number": number, "status": "skipped"})
            continue
        invoice_type = _pick(row, "invoice_type")
        issue_date = _pick(row, "issue_date")
        try:
            from datetime import date as _date

            parsed_date = _date.fromisoformat(issue_date.replace("/", "-")) if issue_date else None
        except ValueError:
            parsed_date = None
        inv = Invoice(
            file_url=f"tenant-default/sales/list-{number}",
            file_type="LIST", invoice_direction="output", invoice_number=number,
            issue_date=parsed_date,
            buyer_name=_pick(row, "buyer_name"), buyer_tax_id=_pick(row, "buyer_tax_id"),
            amount_without_tax=_to_decimal(_pick(row, "amount_without_tax")),
            tax_amount=_to_decimal(_pick(row, "tax_amount")),
            total_amount=_to_decimal(_pick(row, "total_amount")),
            invoice_type=invoice_type, parse_source="SALES_LIST",
            status="parsed", verify_status="pending",
            red_flag=detect_red_invoice(invoice_type, invoice_type),
            user_id=user.id,
        )
        db.add(inv)
        db.flush()
        # 红票：按清单里的"原发票号码"关联蓝票（清单未给 → 留空待人工补）
        _link_original(db, inv, _pick(row, "original_invoice_number"))
        imported += 1
        details.append({"invoice_number": number, "status": "imported",
                        "red_flag": inv.red_flag,
                        "original_invoice_id": inv.original_invoice_id})
    write_audit(
        db, action="SALES_IMPORT", user_id=user.id, channel="web",
        detail={"list": filename, "imported": imported, "skipped": skipped, "errors": errors},
    )
    db.commit()
    return {"imported": imported, "skipped": skipped, "errors": errors, "details": details}


def link_red_invoice(db: Session, user: User, red_id: int, original_id: int) -> Invoice:
    """人工补关联：把红字票挂到原蓝票（仅红票可操作；方向需一致）。"""
    red = db.get(Invoice, red_id)
    original = db.get(Invoice, original_id)
    if red is None or original is None:
        raise ValueError("发票不存在")
    if not red.red_flag:
        raise ValueError("仅红字发票可关联原蓝票")
    if red.invoice_direction != original.invoice_direction:
        raise ValueError("红票与原蓝票方向不一致")
    if red.id == original.id:
        raise ValueError("不能关联自身")
    red.original_invoice_id = original.id
    write_audit(
        db, action="SALES_LINK_RED", user_id=user.id, invoice_id=red.id, channel="web",
        detail={"red_invoice_id": red.id, "original_invoice_id": original.id},
    )
    db.commit()
    return red


def unlinked_red_invoices(db: Session, user: User | None = None) -> list[Invoice]:
    """未关联原蓝票的红字票（财务待处理清单）。"""
    return (
        db.query(Invoice)
        .filter(Invoice.red_flag.is_(True), Invoice.original_invoice_id.is_(None))
        .order_by(Invoice.id.desc())
        .all()
    )
